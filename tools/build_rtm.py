"""Dựng Requirements Traceability Matrix từ tài liệu nguồn và kiểm độ phủ.
BN → US → UC → Function → SCR → TC, kèm AC, BR, flow UCS.
Ghi 08_Documents/RTM.md, RTM.xlsx, UAT-Test-Cases.xlsx. Thoát ≠ 0 nếu có lỗ hổng truy vết."""
import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(__file__).resolve().parent.parent
DOC = ROOT / "08_Documents"
rd = lambda p: (ROOT / p).read_text(encoding="utf-8")
ids = lambda s, pre: re.findall(pre + r"-\d+(?:\.\d+)?", s)


def rows(text, first):
    """Các dòng bảng markdown có ô đầu khớp regex first."""
    out = []
    for line in text.splitlines():
        if line.startswith("|"):
            c = [x.strip() for x in line.strip().strip("|").split("|")]
            if re.fullmatch(first, c[0]):
                out.append(c)
    return out


# --- nguồn ---------------------------------------------------------------------------------
bl = rd("03_User_Story/01_Product-Backlog.md")
US = {c[0]: dict(role=c[1], want=c[2], bn=ids(c[6], "BN"), br=ids(c[7], "BR"), pri=c[8], rel=c[9]) for c in rows(bl, r"US-\d\d")}
AC = defaultdict(list)
cur = None
for line in bl.splitlines():
    m = re.match(r"- \*\*(US-\d\d)\*\*", line)
    if m:
        cur = m[1]
    m = re.search(r"\*\*(AC-\d\d\.\d)\*\*", line)
    if m:
        AC[cur].append(m[1])
UC = {c[0]: dict(name=c[1], us=ids(c[3], "US")) for c in rows(rd("05_Use_Case_Activity/01_Use-Case-Diagram.md"), r"UC-\d\d")}
FN = {c[0]: dict(name=c[1], uc=ids(c[3], "UC"), us=ids(c[4], "US"), scr=ids(c[5], "SCR"))
      for c in rows(rd("06_Function_Feature_State/01_Function-Feature-List.md"), r"F-\d\d\.\d")}
FLOWS = {}
for f in sorted((ROOT / "05_Use_Case_Activity/03_UCS").glob("UCS-*.md")):
    t = f.read_text(encoding="utf-8")
    uc = "UC-" + f.name[4:6]
    FLOWS[uc] = ["BF"] + sorted(set(re.findall(r"^\| ((?:AF|EF)-\d\d) \|", t, re.M)))
uat_text = rd("08_Documents/UAT-Test-Cases.md")
TC_ROWS = rows(uat_text, r"TC-\d{3}")
HEAD = ["TC", "Tên", "Loại", "US · AC", "UC · Flow", "BR · NFR", "Tiền điều kiện", "Bước thực hiện", "Kết quả mong đợi", "Ưu tiên", "Lần 1", "Lần 2"]
TC = {c[0]: dict(name=c[1], us=ids(c[3], "US"), ac=ids(c[3], "AC"), flows=re.findall(r"(UC-\d\d) (BF|AF-\d\d|EF-\d\d)", c[4]),
                 br=ids(c[5], "BR"), nfr=ids(c[5], "NFR"), pri=c[9], r1=c[10], r2=c[11]) for c in TC_ROWS}
BRS = {c[0] for c in rows(rd("05_Use_Case_Activity/02_Business-Rules.md"), r"BR-\d\d")}
NFRS = {c[0] for c in rows(rd("02_Initial_Analysis/05_NFR.md"), r"NFR-\d\d")}
BNS = {c[0]: c[1] for c in rows(rd("01_BACCM_Elicitation/01_BACCM.md"), r"BN-\d\d")}

# --- kiểm phủ ------------------------------------------------------------------------------
gaps = []
mvp = [u for u, d in US.items() if d["rel"] == "MVP"]
uc_of = {u: [k for k, d in UC.items() if u in d["us"]] for u in US}
fn_of = {u: [k for k, d in FN.items() if u in d["us"]] for u in US}
tc_of = {u: [k for k, d in TC.items() if u in d["us"] or any(a in AC[u] for a in d["ac"])] for u in US}
for b in BNS:
    if not any(b in US[u]["bn"] for u in mvp):
        gaps.append(f"{b} không có User Story MVP")
for u in mvp:
    for name, lst in (("Use Case", uc_of[u]), ("Function", fn_of[u]), ("Test Case", tc_of[u])):
        if not lst:
            gaps.append(f"{u} thiếu {name}")
    if not AC[u]:
        gaps.append(f"{u} thiếu Acceptance Criteria")
covered_ac = {a for d in TC.values() for a in d["ac"]}
for u in mvp:
    for a in AC[u]:
        if a not in covered_ac:
            gaps.append(f"{a} chưa có Test Case")
covered_fl = {f"{a} {b}" for d in TC.values() for a, b in d["flows"]}
for uc, fl in FLOWS.items():
    for f in fl:
        if f"{uc} {f}" not in covered_fl:
            gaps.append(f"{uc} {f} chưa có Test Case")
for k, d in UC.items():
    if not d["us"]:
        gaps.append(f"{k} không xuất phát từ User Story")
for d in list(US.values()) + list(TC.values()):
    for b in d["br"]:
        if b not in BRS:
            gaps.append(f"{b} không có trong danh mục Business Rule")
for d in TC.values():
    for n in d["nfr"]:
        if n not in NFRS:
            gaps.append(f"{n} không có trong danh sách NFR")
for b in sorted(BRS):
    if not any(b in d["br"] for d in TC.values()):
        gaps.append(f"{b} chưa được Test Case nào kiểm")

# --- thống kê UAT --------------------------------------------------------------------------
n = len(TC)
high = [k for k, d in TC.items() if d["pri"] == "High"]
r1_pass = sum(d["r1"] == "P" for d in TC.values())
final_pass = sum(d["r1"] == "P" or d["r2"] == "P" for d in TC.values())
print(f"US {len(US)} (MVP {len(mvp)}) · AC {sum(len(AC[u]) for u in mvp)} · UC {len(UC)} · F {len(FN)} · flow UCS {sum(len(v) for v in FLOWS.values())} · TC {n}")
print(f"UAT: High {len(high)} · vòng 1 đạt {r1_pass} ({r1_pass / n:.1%}) · High vòng 1 đạt {sum(TC[k]['r1'] == 'P' for k in high)} · cuối đạt {final_pass}")

# --- RTM.md --------------------------------------------------------------------------------
L = ["# Requirements Traceability Matrix — HL-OMS", "",
     "> Sinh tự động bởi `python tools/build_rtm.py` từ Product Backlog, Use Case list, Function list, UCS và UAT Test Cases. Không sửa tay.", "",
     "## 1. Độ phủ", "", "| Kiểm tra | Kết quả |", "|---|---|",
     f"| Business Need có User Story MVP | {sum(any(b in US[u]['bn'] for u in mvp) for b in BNS)}/{len(BNS)} |",
     f"| User Story MVP có Use Case | {sum(bool(uc_of[u]) for u in mvp)}/{len(mvp)} |",
     f"| User Story MVP có Function | {sum(bool(fn_of[u]) for u in mvp)}/{len(mvp)} |",
     f"| User Story MVP có Test Case | {sum(bool(tc_of[u]) for u in mvp)}/{len(mvp)} |",
     f"| Acceptance Criteria có Test Case | {sum(a in covered_ac for u in mvp for a in AC[u])}/{sum(len(AC[u]) for u in mvp)} |",
     f"| Luồng UCS (BF/AF/EF) có Test Case | {sum(f'{uc} {f}' in covered_fl for uc, fl in FLOWS.items() for f in fl)}/{sum(len(v) for v in FLOWS.values())} |",
     f"| Business Rule được kiểm | {sum(any(b in d['br'] for d in TC.values()) for b in BRS)}/{len(BRS)} |",
     f"| Lỗ hổng truy vết | {len(gaps)} |", "",
     "## 2. Ma trận Business Need → User Story → Use Case → Function → Màn hình → Test Case", "",
     "| BN | US | User Story | UC | Function | SCR | BR | AC | TC |", "|---|---|---|---|---|---|---|---|---|"]
for u in mvp:
    d = US[u]
    scr = sorted({s for f in fn_of[u] for s in FN[f]["scr"]})
    L.append(f"| {', '.join(d['bn'])} | {u} | {d['role']} — {d['want']} | {', '.join(uc_of[u])} | {', '.join(fn_of[u])} | {', '.join(scr)} | "
             f"{', '.join(d['br']) or '—'} | {', '.join(AC[u])} | {', '.join(tc_of[u])} |")
L += ["", "## 3. Business Need → kết quả đo", "", "| BN | Nhu cầu | Số User Story | Số Test Case |", "|---|---|---|---|"]
for b, name in BNS.items():
    us_b = [u for u in mvp if b in US[u]["bn"]]
    L.append(f"| {b} | {name} | {len(us_b)} | {len({t for u in us_b for t in tc_of[u]})} |")
L += ["", "## 4. Luồng UCS → Test Case", "", "| UC | Luồng | Test Case |", "|---|---|---|"]
for uc, fl in FLOWS.items():
    for f in fl:
        L.append(f"| {uc} | {f} | {', '.join(k for k, d in TC.items() if (uc, f) in d['flows'])} |")
L += ["", "## 5. Business Rule → Test Case", "", "| BR | Test Case |", "|---|---|"]
for b in sorted(BRS):
    L.append(f"| {b} | {', '.join(k for k, d in TC.items() if b in d['br'])} |")
if gaps:
    L += ["", "## Lỗ hổng", ""] + [f"- {g}" for g in gaps]
(DOC / "RTM.md").write_text("\n".join(L) + "\n", encoding="utf-8", newline="\n")

# --- Excel ---------------------------------------------------------------------------------
HEADF = PatternFill("solid", fgColor="1F2937")
THIN = Side(style="thin", color="E5E7EB")


def sheet(wb, title, header, data, widths):
    ws = wb.create_sheet(title)
    ws.append(header)
    for r in data:
        ws.append(r)
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, j).column_letter].width = w
        c = ws.cell(1, j)
        c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF", name="Segoe UI"), HEADF, Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Segoe UI", size=10, bold=c.column == 1)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = Border(bottom=THIN)
            if c.value == "P":
                c.fill = PatternFill("solid", fgColor="E3F6EE")
            elif isinstance(c.value, str) and c.value.startswith("F"):
                c.fill = PatternFill("solid", fgColor="FDE8E8")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


wb = Workbook()
wb.remove(wb.active)
sheet(wb, "RTM", ["US", "BN", "User Story", "UC", "Function", "SCR", "BR", "AC", "TC"],
      [[u, ", ".join(US[u]["bn"]), f"{US[u]['role']} — {US[u]['want']}", ", ".join(uc_of[u]), ", ".join(fn_of[u]),
        ", ".join(sorted({s for f in fn_of[u] for s in FN[f]['scr']})), ", ".join(US[u]["br"]), ", ".join(AC[u]), ", ".join(tc_of[u])] for u in mvp],
      [8, 12, 50, 14, 18, 18, 18, 26, 30])
sheet(wb, "Flow → TC", ["UC", "Flow", "TC"], [[uc, f, ", ".join(k for k, d in TC.items() if (uc, f) in d["flows"])] for uc, fl in FLOWS.items() for f in fl], [8, 8, 30])
wb.save(DOC / "RTM.xlsx")
wb = Workbook()
wb.remove(wb.active)
sheet(wb, "UAT Test Cases", HEAD, TC_ROWS, [9, 34, 7, 22, 22, 16, 28, 40, 44, 9, 11, 8])
wb.save(DOC / "UAT-Test-Cases.xlsx")

print(f"Lỗ hổng: {len(gaps)}")
for g in gaps:
    print(" -", g)
sys.exit(1 if gaps else 0)
