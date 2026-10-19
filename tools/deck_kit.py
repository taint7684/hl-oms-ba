"""Bộ dựng slide dùng chung cho HL-OMS (python-pptx): khổ 16:9, nền trắng, một màu nhấn."""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

INK = RGBColor(0x1F, 0x29, 0x37)      # chữ chính
MUTED = RGBColor(0x6B, 0x72, 0x80)    # chữ phụ
LINE = RGBColor(0xE5, 0xE7, 0xEB)     # đường kẻ
SOFT = RGBColor(0xF6, 0xF4, 0xF1)     # nền khối
ACCENT = RGBColor(0xB4, 0x45, 0x1F)   # màu nhấn (đất nung)
ACCENT_SOFT = RGBColor(0xFB, 0xEE, 0xE8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Segoe UI"

W, H = Inches(13.333), Inches(7.5)
MX = Inches(0.6)                      # lề ngang
CW = W - 2 * MX                       # bề rộng nội dung


def new_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    return prs


def _i(*v):
    """Tọa độ phải là số nguyên EMU (phép chia ra float làm PowerPoint không mở được file)."""
    return [Emu(int(round(x))) for x in v]


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def text(slide, x, y, w, h, value, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font=FONT):
    """Hộp chữ; value là str hoặc list dòng (dòng bắt đầu '• ' thành gạch đầu dòng)."""
    x, y, w, h = _i(x, y, w, h)
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    lines = value if isinstance(value, list) else [value]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(size * 0.45)
        r = p.add_run()
        r.text = line
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, font
        r.font.color.rgb = color
    return tb


def rect(slide, x, y, w, h, fill=SOFT, line=None, shape=MSO_SHAPE.RECTANGLE):
    x, y, w, h = _i(x, y, w, h)
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    return s


def header(prs, kicker, title, page, total):
    """Slide nội dung: nhãn nhỏ + tiêu đề + chân trang."""
    s = _blank(prs)
    rect(s, MX, Inches(0.55), Inches(0.08), Inches(0.75), fill=ACCENT)
    text(s, MX + Inches(0.22), Inches(0.5), CW, Inches(0.35), kicker.upper(), 11, True, ACCENT)
    text(s, MX + Inches(0.22), Inches(0.82), CW, Inches(0.6), title, 26, True)
    s.shapes.add_connector(1, MX, H - Inches(0.55), W - MX, H - Inches(0.55)).line.color.rgb = LINE
    text(s, MX, H - Inches(0.5), Inches(6), Inches(0.3), "HL-OMS · Họp khởi động · 03/11/2026", 9, color=MUTED)
    text(s, W - MX - Inches(2), H - Inches(0.5), Inches(2), Inches(0.3), f"{page} / {total}", 9,
         color=MUTED, align=PP_ALIGN.RIGHT)
    return s


def table(slide, x, y, w, rows, col_w=None, size=11, head_size=11, row_h=Inches(0.42)):
    """Bảng: rows[0] là tiêu đề. col_w là tỉ lệ cột."""
    x, y, w, row_h = _i(x, y, w, row_h)
    nr, nc = len(rows), len(rows[0])
    shp = slide.shapes.add_table(nr, nc, x, y, w, row_h * nr)
    tbl = shp.table
    tbl.first_row = True
    ratios = col_w or [1] * nc
    tot = sum(ratios)
    for j, r in enumerate(ratios):
        tbl.columns[j].width = Emu(int(w * r / tot))
    for i in range(nr):
        tbl.rows[i].height = row_h
        for j in range(nc):
            c = tbl.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = INK if i == 0 else (WHITE if i % 2 else SOFT)
            c.margin_left = c.margin_right = Inches(0.1)
            c.margin_top = c.margin_bottom = Inches(0.04)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = ""
            r = p.add_run()
            r.text = str(rows[i][j])
            r.font.name = FONT
            r.font.size = Pt(head_size if i == 0 else size)
            r.font.bold = i == 0 or j == 0
            r.font.color.rgb = WHITE if i == 0 else INK
    return shp


def card(slide, x, y, w, h, title, body, size=12, fill=SOFT, title_color=ACCENT):
    rect(slide, x, y, w, h, fill=fill)
    text(slide, x + Inches(0.2), y + Inches(0.15), w - Inches(0.4), Inches(0.4), title, 14, True, title_color)
    text(slide, x + Inches(0.2), y + Inches(0.6), w - Inches(0.4), h - Inches(0.7), body, size)
