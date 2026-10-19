"""Xuất BRD, SRS, UAT, User Guide, Go-live Plan (.md là bản nguồn) ra Word trong 08_Documents/word/.
Gộp các dòng `<!-- include: path -->`, vẽ khối mermaid thành PNG, dàn bảng và ảnh theo khổ trang."""
import re
import subprocess
import sys
from pathlib import Path

import pypandoc
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "08_Documents"
OUT = DOCS / "word"
BUILD = ROOT / ".build"
TARGETS = [  # (nguồn, file Word, khổ ngang?)
    ("BRD.md", "HL-OMS_BRD.docx", False),
    ("SRS.md", "HL-OMS_SRS.docx", False),
    ("UAT-Test-Cases.md", "HL-OMS_UAT.docx", True),
    ("User-Guide.md", "HL-OMS_User-Guide.docx", False),
    ("Go-Live-Plan.md", "HL-OMS_Go-Live-Plan.docx", False),
]
INK, ACC, MUTED = RGBColor(0x1F, 0x29, 0x37), RGBColor(0xB4, 0x45, 0x1F), RGBColor(0x6B, 0x72, 0x80)
FONT = "Segoe UI"


def rebase(text, src_dir):
    """Ảnh → đường dẫn tuyệt đối; link tới tài liệu khác → chỉ giữ chữ."""
    def img(m):
        p = (src_dir / m[2]).resolve()
        return f"![{m[1]}]({p.as_posix()})"
    text = re.sub(r"!\[([^\]]*)\]\(([^)]+\.(?:png|jpg|svg))\)", img, text)
    text = re.sub(r"(?<!!)\[([^\]]+)\]\((?!https?:)[^)]+\)", r"\1", text)
    return text


def expand(path, depth=0):
    out = []
    for line in path.read_text(encoding="utf-8").splitlines():
        m = re.search(r"<!-- include: ([^ ]+) -->", line)
        if m and depth < 2:
            inc = (path.parent / m[1]).resolve()
            body = expand(inc, depth + 1)
            body = re.sub(r"^(#+) ", lambda h: "#" * min(6, len(h[1]) + 3) + " ", body, flags=re.M)
            out.append(body)
        else:
            out.append(rebase(line, path.parent))
    return "\n".join(out)


def mermaid_blocks(text, key):
    blocks = re.findall(r"```mermaid\n(.*?)```", text, flags=re.S)
    if not blocks:
        return text
    BUILD.mkdir(exist_ok=True)
    files = []
    for i, code in enumerate(blocks):
        f = BUILD / f"{key}_{i}.mmd"
        f.write_text(code, encoding="utf-8", newline="\n")
        files.append(f)
    subprocess.run([sys.executable, str(ROOT / "tools/mermaid.py"), *map(str, files)], check=True, capture_output=True)
    it = iter(files)
    return re.sub(r"```mermaid\n.*?```", lambda m: f"![]({next(it).with_suffix('.png').as_posix()})", text, flags=re.S)


def reference_doc(landscape):
    ref = BUILD / f"ref_{'L' if landscape else 'P'}.docx"
    raw = subprocess.run([pypandoc.get_pandoc_path(), "--print-default-data-file", "reference.docx"], capture_output=True, check=True).stdout
    ref.write_bytes(raw)
    d = Document(ref)
    st = d.styles
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if name in [s.name for s in st]:
            s = st[name]
            s.font.name, s.font.size, s.font.color.rgb = FONT, Pt(10), INK
            s.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
            s.paragraph_format.space_after = Pt(4)
    for lvl, size in ((1, 18), (2, 14), (3, 12), (4, 11), (5, 10.5), (6, 10)):
        s = st[f"Heading {lvl}"]
        s.font.name, s.font.size, s.font.bold = FONT, Pt(size), True
        s.font.color.rgb = ACC if lvl in (2,) else INK
        s.paragraph_format.space_before = Pt(14 if lvl <= 2 else 10)
        s.paragraph_format.space_after = Pt(6)
        rpr = s.element.get_or_add_rPr()
        for el in rpr.findall(qn("w:rFonts")):
            rpr.remove(el)
        f = OxmlElement("w:rFonts")
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            f.set(qn(a), FONT)
        rpr.append(f)
    sec = d.sections[0]
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Cm(29.7), Cm(21)
    else:
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(1.8))
    sec.top_margin = sec.bottom_margin = Cm(1.7)
    d.save(ref)
    return ref


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"), shd.set(qn("w:color"), "auto"), shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom", "insideH", "left", "right", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single" if edge in ("top", "bottom", "insideH") else "nil")
        e.set(qn("w:sz"), "4"), e.set(qn("w:color"), "D1D5DB")
        b.append(e)
    tblPr.append(b)


def page_field(par):
    r = par.add_run()
    for t, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if t:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), t)
            r._r.append(fc)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = txt
            r._r.append(it)


def polish(path, title, landscape):
    d = Document(path)
    sec = d.sections[0]
    width = sec.page_width - sec.left_margin - sec.right_margin
    for t in d.tables:
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        borders(t)
        tw = t._tbl.tblPr.find(qn("w:tblW"))                 # bảng trải hết bề rộng trang
        if tw is None:
            tw = OxmlElement("w:tblW")
            t._tbl.tblPr.append(tw)
        tw.set(qn("w:type"), "pct"), tw.set(qn("w:w"), "5000")
        trPr = t.rows[0]._tr.get_or_add_trPr()                # lặp dòng tiêu đề khi sang trang
        h = OxmlElement("w:tblHeader")
        h.set(qn("w:val"), "true")
        trPr.append(h)
        for i, row in enumerate(t.rows):
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.space_after = Pt(1)
                    for r in p.runs:
                        r.font.size = Pt(8 if landscape else 8.5)
                        if i == 0:
                            r.font.bold, r.font.color.rgb = True, RGBColor(0xFF, 0xFF, 0xFF)
                if i == 0:
                    shade(c, "1F2937")
                elif i % 2 == 0:
                    shade(c, "F6F4F1")
    for shape in d.inline_shapes:                     # ảnh không vượt khổ trang / ô bảng
        in_cell = shape._inline.getparent().getparent().getparent().getparent().tag.endswith("tc")
        maxw = int(width * (0.32 if in_cell else 1))
        if shape.width > maxw:
            ratio = maxw / shape.width
            shape.width, shape.height = maxw, int(shape.height * ratio)
    fp = sec.footer.paragraphs[0]
    fp.text = ""
    run = fp.add_run(f"HL-OMS · {title} · Trang ")
    run.font.size, run.font.color.rgb = Pt(8), MUTED
    page_field(fp)
    d.core_properties.title = f"HL-OMS {title}"
    d.core_properties.author = "Business Analyst"
    d.save(path)


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    BUILD.mkdir(exist_ok=True)
    refs = {False: reference_doc(False), True: reference_doc(True)}
    for src, out, land in TARGETS:
        text = mermaid_blocks(expand(DOCS / src), Path(src).stem)
        text = re.sub(r"^> ", "", text, flags=re.M)
        target = OUT / out
        pypandoc.convert_text(text, "docx", format="gfm+pipe_tables", outputfile=str(target),
                              extra_args=[f"--reference-doc={refs[land]}", f"--resource-path={DOCS}"])
        polish(target, Path(out).stem.replace("HL-OMS_", ""), land)
        print(f"{out}: {target.stat().st_size // 1024} KB")
