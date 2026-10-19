"""Xuất các log ở 00_Admin (.md là bản nguồn) ra một file Excel: 00_Admin/HL-OMS_Registers.xlsx."""
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(__file__).resolve().parent.parent
SHEETS = [  # (tên sheet, file nguồn, cột tối thiểu để nhận là bảng dữ liệu)
    ("Assumption Log", "00_Admin/Assumption-Log.md", 6),
    ("Question Log", "00_Admin/Question-Log.md", 8),
    ("Decision Log", "00_Admin/Decision-Log.md", 6),
    ("Change Log", "00_Admin/Change-Log.md", 6),
]
HEAD = PatternFill("solid", fgColor="1F2937")
ZEBRA = PatternFill("solid", fgColor="F6F4F1")
THIN = Side(style="thin", color="E5E7EB")


def clean(cell):
    cell = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", cell)   # bỏ link markdown, giữ chữ
    return cell.replace("**", "").replace("`", "").strip()


def tables(md, min_cols):
    """Trả (tiêu đề, các dòng) gộp mọi bảng cùng tiêu đề; thêm cột Nhóm theo heading ### gần nhất."""
    header, rows, group = None, [], ""
    lines = md.splitlines()
    for i, line in enumerate(lines):
        if line.startswith("### "):
            group = line[4:].strip()
        if not line.startswith("|"):
            continue
        cells = [clean(c) for c in line.strip().strip("|").split("|")]
        nxt = lines[i + 1].strip() if i + 1 < len(lines) else ""
        if re.match(r"^\|[\s:|-]+\|$", nxt) and len(cells) >= min_cols:
            if header is None:
                header = (["Nhóm"] if group else []) + cells
            continue
        if re.match(r"^\|[\s:|-]+\|$", line.strip()) or header is None:
            continue
        if len(cells) >= min_cols and not cells[0].startswith("Tổng"):
            if (["Nhóm"] if group else []) + cells != header:
                rows.append(([group] if group else []) + cells)
    return header, rows


wb = Workbook()
wb.remove(wb.active)
for name, src, min_cols in SHEETS:
    header, rows = tables((ROOT / src).read_text(encoding="utf-8"), min_cols)
    ws = wb.create_sheet(name)
    ws.append(header)
    for r in rows:
        ws.append(r[: len(header)])
    for j, _ in enumerate(header, 1):
        c = ws.cell(1, j)
        c.font, c.fill = Font(bold=True, color="FFFFFF", name="Segoe UI"), HEAD
        c.alignment = Alignment(vertical="center", wrap_text=True)
        width = max(len(str(ws.cell(i, j).value or "")) for i in range(1, ws.max_row + 1))
        ws.column_dimensions[c.column_letter].width = min(max(width * 0.9, 8), 60)
    for i in range(2, ws.max_row + 1):
        for j in range(1, len(header) + 1):
            c = ws.cell(i, j)
            c.font = Font(name="Segoe UI", size=10, bold=(j == 1 and header[0] != "Nhóm") or header[j - 1] == "ID")
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = Border(bottom=THIN)
            if i % 2 == 0:
                c.fill = ZEBRA
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    print(f"{name}: {len(rows)} dòng")

out = ROOT / "00_Admin/HL-OMS_Registers.xlsx"
wb.save(out)
print(f"Đã ghi {out}")
